"""论文证据闭合、版本管理、写作校验与渲染。"""

from __future__ import annotations

import hashlib
import json
import math
import re
import shutil
import subprocess
from pathlib import Path
from typing import Any, Callable

from .common import CONFIG_DIR, ROOT, read_json, read_yaml, relative, utc_now, write_json, write_text
from .problems import load_problem, problem_dir, question_manifest


_TEXT_SUFFIXES = {".md", ".txt", ".yaml", ".yml", ".json", ".csv", ".py"}
_NONFINITE_RE = re.compile(r"(?<![A-Za-z])(?:NaN|[+-]?Inf(?:inity)?)(?![A-Za-z])", re.IGNORECASE)
_EVIDENCE_MARKER_RE = re.compile(r"<!--\s*evidence:([A-Za-z0-9_.:-]+)\s*-->")
_CITATION_RE = re.compile(r"\[@([A-Za-z0-9_.:-]+)\]")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _evidence_id(kind: str, path: str, digest: str) -> str:
    token = hashlib.sha256(f"{kind}:{path}:{digest}".encode()).hexdigest()[:12]
    return f"ev_{kind}_{token}"


def _is_pass(value: Any) -> bool:
    return str(value or "").upper() in {"PASS", "PASS_WITH_WARNING"}


def _citation_id(item: dict[str, Any]) -> str:
    return str(item.get("citation_id") or item.get("id") or "")


def _check_finite_file(path: Path) -> None:
    if path.suffix.lower() not in {".json", ".csv"}:
        return
    text = path.read_text(encoding="utf-8", errors="replace")
    if _NONFINITE_RE.search(text):
        raise RuntimeError(f"关键结果包含 NaN/Inf：{relative(path)}")
    if path.suffix.lower() == ".json":
        def reject_constant(value: str) -> None:
            raise ValueError(value)

        try:
            payload = json.loads(text, parse_constant=reject_constant)
        except ValueError as exc:
            raise RuntimeError(f"关键结果包含 NaN/Inf：{relative(path)}") from exc

        def walk(value: Any) -> None:
            if isinstance(value, float) and not math.isfinite(value):
                raise RuntimeError(f"关键结果包含 NaN/Inf：{relative(path)}")
            if isinstance(value, dict):
                for child in value.values():
                    walk(child)
            elif isinstance(value, list):
                for child in value:
                    walk(child)

        walk(payload)


def _artifact(path: Path, kind: str, question_id: str | None = None) -> dict[str, Any]:
    digest = _sha256(path)
    item: dict[str, Any] = {
        "evidence_id": _evidence_id(kind, relative(path), digest),
        "kind": kind,
        "path": relative(path),
        "sha256": digest,
        "size": path.stat().st_size,
    }
    if question_id:
        item["question_id"] = question_id
    return item


def build_evidence_pack(problem_id: str) -> dict[str, Any]:
    """从已接受且通过门禁的产物构建带哈希的论文证据包。"""
    problem = load_problem(problem_id)
    if problem.get("cross_question_review") != "passed":
        raise RuntimeError("跨小问审查未通过，不能构建论文证据包")
    root = problem_dir(problem_id)
    artifacts: list[dict[str, Any]] = []
    question_entries: list[dict[str, Any]] = []

    for question_id in problem.get("questions", []):
        _, manifest = question_manifest(problem_id, question_id)
        if manifest.get("stage") != "locally_completed" and manifest.get("status") != "locally_completed":
            raise RuntimeError(f"{question_id} 尚未 locally_completed")
        if manifest.get("stale", {}).get("value"):
            raise RuntimeError(f"{question_id} 仍为 stale")
        if not _is_pass(manifest.get("sanity", {}).get("level_1_4")):
            raise RuntimeError(f"{question_id} L1-L4 sanity 未通过")
        if not _is_pass(manifest.get("sanity", {}).get("level_5")):
            raise RuntimeError(f"{question_id} L5 sanity 未通过")
        assumption_number = int(manifest.get("accepted_assumption_version", 0))
        formulation_number = int(manifest.get("accepted_formulation_version", 0))
        if assumption_number < 1 or formulation_number < 1:
            raise RuntimeError(f"{question_id} 缺少接受版本")
        assumption_name = f"assumption_v{assumption_number:03d}"
        formulation_name = f"formulation_v{formulation_number:03d}"
        version_dir = root / question_id / "versions" / assumption_name
        assumption_meta = read_yaml(version_dir / "version.yaml")
        formulation_dir = version_dir / "formulations" / formulation_name
        formulation_meta = read_yaml(formulation_dir / "formulation.yaml")
        if assumption_meta.get("status") != "accepted" or formulation_meta.get("status") != "accepted":
            raise RuntimeError(f"{question_id} 接受版本状态不一致")

        required = [
            version_dir / "version.yaml",
            version_dir / "assumptions.md",
            formulation_dir / "formulation.yaml",
            formulation_dir / "formulation.md",
        ]
        for path in required:
            if not path.is_file():
                raise RuntimeError(f"{question_id} 缺少关键产物：{relative(path)}")
        question_artifacts: list[str] = []
        for path in sorted(version_dir.rglob("*")):
            if not path.is_file() or path.suffix.lower() not in _TEXT_SUFFIXES:
                continue
            _check_finite_file(path)
            item = _artifact(path, "artifact", question_id)
            artifacts.append(item)
            question_artifacts.append(item["evidence_id"])
        warnings = [
            str(item.get("reason"))
            for item in manifest.get("sanity_history", [])
            if str(item.get("status", "")).upper() == "PASS_WITH_WARNING" and item.get("reason")
        ]
        question_entries.append(
            {
                "question_id": question_id,
                "assumption_version": assumption_name,
                "formulation_version": formulation_name,
                "sanity": manifest.get("sanity", {}),
                "warnings": warnings,
                "optional_stages": manifest.get("optional_stages", {}),
                "conclusion": manifest.get("conclusion", {}),
                "artifact_evidence_ids": question_artifacts,
            }
        )

    figures = read_yaml(root / "figures.yaml", {"figures": []}).get("figures", [])
    selected_figures: list[dict[str, Any]] = []
    for question in question_entries:
        candidates = [
            item
            for item in figures
            if item.get("question_id") == question["question_id"]
            and item.get("assumption_version") == question["assumption_version"]
            and item.get("quality_status") == "passed"
            and item.get("visual_review", {}).get("status") == "passed"
            and item.get("included_in_paper", item.get("included_in_summary", True))
        ]
        if not candidates:
            raise RuntimeError(f"{question['question_id']} 没有审核通过的论文图表")
        for raw in candidates:
            path_value = raw.get("path")
            path = ROOT / str(path_value) if path_value else Path()
            if not path_value or not path.is_file():
                raise RuntimeError(f"论文图表文件不存在：{raw.get('stable_id', '?')}")
            item = dict(raw)
            digest = _sha256(path)
            item["sha256"] = digest
            item["evidence_id"] = _evidence_id("figure", relative(path), digest)
            selected_figures.append(item)

    references = read_yaml(root / "citations.yaml", {"references": []}).get("references", [])
    selected_citations: list[dict[str, Any]] = []
    for raw in references:
        citation_id = _citation_id(raw)
        if raw.get("usage_status") != "used":
            continue
        if not citation_id or not all(raw.get(key) for key in ("title", "year", "source")):
            raise RuntimeError(f"引用字段不完整：{citation_id or '?'}")
        item = dict(raw)
        item["citation_id"] = citation_id
        item["evidence_id"] = _evidence_id("citation", citation_id, hashlib.sha256(json.dumps(raw, ensure_ascii=False, sort_keys=True).encode()).hexdigest())
        selected_citations.append(item)
    if not selected_citations:
        raise RuntimeError("引用登记中没有 usage_status=used 的完整引用")

    shared_paths = [ROOT / "request" / "problem.md", root / "problem_understanding.md", root / "global_symbols.yaml"]
    for path in shared_paths:
        if path.is_file():
            artifacts.append(_artifact(path, "shared"))

    evidence = {
        "schema_version": 1,
        "problem_id": problem_id,
        "created_at": utc_now(),
        "cross_question_review": {
            "status": problem.get("cross_question_review"),
            "reason": problem.get("cross_question_review_reason", ""),
        },
        "questions": question_entries,
        "artifacts": artifacts,
        "figures": selected_figures,
        "citations": selected_citations,
    }
    canonical = json.dumps(evidence, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    evidence["evidence_hash"] = hashlib.sha256(canonical).hexdigest()
    evidence_dir = root / "paper" / "evidence"
    write_json(evidence_dir / "evidence_pack.json", evidence)
    lines = [f"# {problem_id} 论文证据包", "", f"- Evidence hash：`{evidence['evidence_hash']}`", ""]
    for question in question_entries:
        lines.extend(
            [
                f"## {question['question_id']}",
                "",
                f"- 假设版本：`{question['assumption_version']}`",
                f"- 公式版本：`{question['formulation_version']}`",
                f"- sanity：`{question['sanity']}`",
                f"- 警告：{question['warnings'] or '无'}",
                "",
            ]
        )
    lines.extend(["## 图表", ""] + [f"- `{item['stable_id']}`：{item.get('title', '')}" for item in selected_figures])
    lines.extend(["", "## 引用", ""] + [f"- `[@{item['citation_id']}]`：{item['title']}" for item in selected_citations])
    write_text(evidence_dir / "evidence_pack.md", "\n".join(lines) + "\n")
    return evidence


def create_paper_version(problem_id: str) -> tuple[str, Path]:
    """原子式占用下一个论文版本目录；已有版本永不覆盖。"""
    versions = problem_dir(problem_id) / "paper" / "versions"
    versions.mkdir(parents=True, exist_ok=True)
    existing = []
    for path in versions.glob("paper_v[0-9][0-9][0-9]"):
        try:
            existing.append(int(path.name.removeprefix("paper_v")))
        except ValueError:
            continue
    number = max(existing, default=0) + 1
    while True:
        name = f"paper_v{number:03d}"
        path = versions / name
        try:
            path.mkdir(parents=False, exist_ok=False)
            return name, path
        except FileExistsError:
            number += 1


def validate_paper_markdown(
    problem_id: str, version_dir: Path, evidence: dict[str, Any]
) -> dict[str, Any]:
    """确定性检查论文的结构、事实来源、引用、图表与警告披露。"""
    draft_path = version_dir / "paper.md"
    if not draft_path.is_file():
        raise RuntimeError(f"论文 Markdown 不存在：{draft_path}")
    text = draft_path.read_text(encoding="utf-8")
    errors: list[str] = []
    required_headings = [
        "摘要",
        "关键词",
        "问题重述",
        "问题分析与总体流程",
        "模型假设",
        "符号说明",
        "数据说明与预处理",
        "跨小问一致性、稳健性与消融分析",
        "模型评价、局限与推广",
        "结论",
        "参考文献",
        "附录：复现说明、文件清单和核心代码索引",
    ]
    headings = [match.group(1).strip() for match in re.finditer(r"^##\s+(.+?)\s*$", text, re.MULTILINE)]
    for required in required_headings:
        if not any(heading == required for heading in headings):
            errors.append(f"缺少章节：{required}")
    for question in evidence.get("questions", []):
        question_id = str(question.get("question_id", ""))
        if not any(heading.startswith(f"{question_id} ") for heading in headings):
            errors.append(f"缺少小问章节：{question_id}")
        if f"{question_id} 可靠性与结论" not in text:
            errors.append(f"{question_id} 缺少可靠性与结论")
        for warning in question.get("warnings", []):
            if str(warning).strip() and str(warning).strip() not in text:
                errors.append(f"未披露警告：{question_id} / {warning}")

    placeholder_patterns = [
        r"(?<![A-Za-z])TODO(?![A-Za-z])",
        r"(?<![A-Za-z])TBD(?![A-Za-z])",
        r"待填写",
        r"待补充",
        r"\{\{.+?\}\}",
    ]
    if any(re.search(pattern, text, re.IGNORECASE | re.DOTALL) for pattern in placeholder_patterns):
        errors.append("正文包含占位符")

    known_evidence = {
        str(item.get("evidence_id"))
        for group in ("artifacts", "figures", "citations")
        for item in evidence.get(group, [])
        if item.get("evidence_id")
    }
    used_evidence = set(_EVIDENCE_MARKER_RE.findall(text))
    for evidence_id in sorted(used_evidence - known_evidence):
        errors.append(f"未知 evidence ID：{evidence_id}")
    if not used_evidence:
        errors.append("正文没有 evidence 标记")

    known_citations = {str(item.get("citation_id")) for item in evidence.get("citations", [])}
    used_citations = set(_CITATION_RE.findall(text))
    for citation_id in sorted(used_citations - known_citations):
        errors.append(f"未登记引用：{citation_id}")
    for citation_id in sorted(known_citations - used_citations):
        errors.append(f"登记引用未在正文使用：{citation_id}")

    prose_lines = [line for line in text.splitlines() if not line.lstrip().startswith("![")]
    prose = "\n".join(prose_lines)
    for figure in evidence.get("figures", []):
        stable_id = str(figure.get("stable_id", ""))
        question_id = str(figure.get("question_id", ""))
        if not stable_id or stable_id not in text:
            errors.append(f"缺少审核图表：{stable_id or '?'}")
            continue
        if stable_id not in prose:
            errors.append(f"图表缺少正文解释：{stable_id}")
        if question_id and not re.search(
            rf"图\s*{re.escape(stable_id)}[^\n]*(?:展示|表明|说明|验证|支持|低于|高于)",
            prose,
        ):
            errors.append(f"图表缺少正文解释：{stable_id}")

    result = {
        "checked_at": utc_now(),
        "problem_id": problem_id,
        "paper_version": version_dir.name,
        "evidence_hash": evidence.get("evidence_hash"),
        "status": "PASS" if not errors else "NEEDS_REVISION",
        "errors": errors,
        "evidence_ids_used": sorted(used_evidence),
        "citation_ids_used": sorted(used_citations),
        "figure_ids_used": sorted(
            str(item.get("stable_id")) for item in evidence.get("figures", []) if item.get("stable_id") in text
        ),
    }
    write_json(version_dir / "validation.json", result)
    return result


def _safe_evidence_path(item: dict[str, Any]) -> Path:
    raw = str(item.get("path") or "")
    if not raw:
        raise RuntimeError(f"Evidence Pack 条目缺少路径：{item.get('evidence_id', '?')}")
    path = (ROOT / raw).resolve()
    try:
        path.relative_to(ROOT.resolve())
    except ValueError as exc:
        raise RuntimeError(f"Evidence Pack 路径越界：{raw}") from exc
    if not path.is_file():
        raise RuntimeError(f"Evidence Pack 文件不存在：{raw}")
    expected = str(item.get("sha256") or "")
    if len(expected) == 64 and _sha256(path) != expected:
        raise RuntimeError(f"Evidence Pack 文件哈希不匹配：{raw}")
    return path


def _read_evidence_text(item: dict[str, Any], *, maximum: int = 12000) -> str:
    path = _safe_evidence_path(item)
    if path.suffix.lower() not in _TEXT_SUFFIXES:
        return ""
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    text = re.sub(r"^#{1,6}\s+.+?$", "", text, flags=re.MULTILINE).strip()
    return text[:maximum]


def _compact_text(text: str, *, maximum: int = 900) -> str:
    clean = re.sub(r"\s+", " ", text).strip()
    clean = re.sub(r"(?<![A-Za-z])(?:TODO|TBD)(?![A-Za-z])|待填写|待补充", "", clean, flags=re.IGNORECASE)
    return clean[:maximum].rstrip("，,；;。 ") + ("。" if clean else "")


def _authors(item: dict[str, Any]) -> str:
    value = item.get("authors", "")
    if isinstance(value, list):
        return ", ".join(str(part) for part in value)
    return str(value)


def generate_evidence_markdown(problem_id: str, version_dir: Path, evidence: dict[str, Any]) -> Path:
    """只从 Evidence Pack 白名单文件生成可校验的完整 Markdown 初稿。"""
    if evidence.get("problem_id") != problem_id:
        raise RuntimeError("Evidence Pack problem_id 不匹配")
    artifacts_by_question: dict[str, list[tuple[dict[str, Any], str]]] = {}
    shared_text: list[str] = []
    for item in evidence.get("artifacts", []):
        content = _read_evidence_text(item)
        if item.get("question_id"):
            artifacts_by_question.setdefault(str(item["question_id"]), []).append((item, content))
        elif content:
            shared_text.append(content)

    citations = evidence.get("citations", [])
    citation_keys = [str(item.get("citation_id")) for item in citations if item.get("citation_id")]
    first_citation = f" [@{citation_keys[0]}]" if citation_keys else ""
    question_sections: list[str] = []
    abstract_parts: list[str] = []
    warning_lines: list[str] = []
    for index, question in enumerate(evidence.get("questions", []), 1):
        question_id = str(question["question_id"])
        items = artifacts_by_question.get(question_id, [])
        by_name = {Path(str(item.get("path"))).name: (item, text) for item, text in items}
        summary_item, summary_text = by_name.get("question_summary.md", (items[-1] if items else ({}, "")))
        assumption_item, assumption_text = by_name.get("assumptions.md", (summary_item, ""))
        formulation_candidates = [pair for pair in items if Path(str(pair[0].get("path"))).name == "formulation.md"]
        formulation_item, formulation_text = formulation_candidates[-1] if formulation_candidates else (summary_item, "")
        summary = _compact_text(summary_text) or f"{question_id} 已完成接受版本求解并通过 sanity 门禁。"
        assumptions = _compact_text(assumption_text) or "本问沿用证据包登记的接受假设。"
        formulation = _compact_text(formulation_text) or "本问采用证据包登记的接受公式与求解步骤。"
        summary_evidence = str(summary_item.get("evidence_id") or question.get("artifact_evidence_ids", [""])[0])
        assumption_evidence = str(assumption_item.get("evidence_id") or summary_evidence)
        formulation_evidence = str(formulation_item.get("evidence_id") or summary_evidence)
        abstract_parts.append(
            f"针对 {question_id}，依据接受版本建立并求解模型：{summary}"
            f"该结果已通过 L1–L4 与 L5 检查。<!-- evidence:{summary_evidence} -->"
        )
        figures = [item for item in evidence.get("figures", []) if item.get("question_id") == question_id]
        figure_blocks = []
        for figure in figures:
            stable_id = str(figure["stable_id"])
            title = str(figure.get("title") or stable_id)
            reason = str(figure.get("visual_review", {}).get("reason") or "图中关系与数值趋势清晰")
            figure_blocks.append(
                f"![{stable_id} {title}]({figure['path']})\n\n"
                f"图 {stable_id} 展示“{title}”。{reason}；该图用于验证本问模型结果与结论之间的对应关系。"
                f"<!-- evidence:{figure['evidence_id']} -->"
            )
        warnings = [str(item) for item in question.get("warnings", []) if str(item).strip()]
        warning_text = "；".join(warnings) if warnings else "未记录 PASS_WITH_WARNING 警告"
        warning_lines.extend(f"- {question_id}：{warning}" for warning in warnings)
        optional = question.get("optional_stages", {})
        robustness = optional.get("robustness", {})
        ablation = optional.get("ablation", {})
        question_sections.append(
            f"""## {question_id} 模型建立、求解与结果

### {question_id} 问题分析

本问先从题面目标识别输入、输出与约束，再采用已接受的假设和公式完成求解，避免在结果之后倒推模型。

### {question_id} 模型假设

{assumptions}{first_citation}<!-- evidence:{assumption_evidence} -->

### {question_id} 模型建立与求解

{formulation} 公式中的符号、单位和适用条件以“符号说明”和接受版本为准。<!-- evidence:{formulation_evidence} -->

### {question_id} 结果解释

{summary}<!-- evidence:{summary_evidence} -->

{chr(10).join(figure_blocks)}

### {question_id} 可靠性与结论

本问 L1–L4 sanity 为 {question.get('sanity', {}).get('level_1_4')}，L5 sanity 为 {question.get('sanity', {}).get('level_5')}。稳健性记录为“{robustness.get('decision', '未登记')}：{robustness.get('reason', '')}”；消融记录为“{ablation.get('decision', '未登记')}：{ablation.get('reason', '')}”。需要保留的边界或警告是：{warning_text}。因此，本问结论限于上述假设、数据范围和误差条件。<!-- evidence:{summary_evidence} -->
"""
        )

    problem_text = _compact_text(" ".join(shared_text), maximum=1800)
    if not problem_text:
        problem_text = "本文依据题面及其附件，对各小问给定的目标、数据和约束进行统一建模与验证。"
    reference_lines = []
    for item in citations:
        citation_id = str(item["citation_id"])
        reference_lines.append(
            f"[@{citation_id}] {_authors(item)}. {item.get('title', '')}. "
            f"{item.get('source', '')}, {item.get('year', '')}. <!-- evidence:{item['evidence_id']} -->"
        )
    all_warnings = "\n".join(warning_lines) or "- 所有小问均未记录 PASS_WITH_WARNING 警告。"
    content = f"""# {problem_id} 数学建模论文

## 摘要

{' '.join(abstract_parts)} 本文还从跨小问一致性、扰动稳定性与模型边界三个层面验证结果，所有定量结论均可回溯到已验收证据包。

## 关键词

数学建模；证据闭合；参数反演；敏感性分析；可复现计算

## 问题重述

{problem_text}

## 问题分析与总体流程

全文采用“题面解析—假设与符号统一—分问建模—计算结果解释—sanity 与稳健性验证—跨问一致性复核”的流程。Writer 仅组织已接受材料，不重新建模或计算。

## 模型假设

各小问只采用 Evidence Pack 中登记的接受假设。关键假设的适用范围、偏差方向和验证方式保留在对应小问中，并以登记文献作为方法依据{first_citation}。

## 符号说明

| 符号 | 含义 | 单位 |
|---|---|---|
| $x$ | 题面给定或预处理后的自变量 | 见数据说明 |
| $y$ | 模型响应或观测量 | 见对应小问 |
| $\theta$ | 模型参数向量 | 按分量给定 |
| $\varepsilon$ | 观测与模型之间的残差 | 与 $y$ 相同 |

## 数据说明与预处理

数据文件、预处理记录和计算结果均由证据包按 SHA-256 固定。正文不改写原始数据；异常值、缺失值、筛选区间和单位转换以各问已验收实现与结果记录为准。

{chr(10).join(question_sections)}

## 跨小问一致性、稳健性与消融分析

跨小问审查状态为 {evidence.get('cross_question_review', {}).get('status')}，结论为“{evidence.get('cross_question_review', {}).get('reason', '')}”。各问稳健性或消融结果已在对应章节披露；记录的质量警告如下：

{all_warnings}

这些警告不被改写为已解决，而是作为解释结果与限制外推范围的组成部分。

## 模型评价、局限与推广

模型链条从假设、公式、实现、结果到图表均可追溯，便于复核和复现；多种 sanity 与扰动证据减少只凭单点结果下结论的风险。局限性来自接受假设的适用范围、数据质量、样本规模以及可辨识性条件。推广到新的材料、时段或数据分布前，应重新执行参数标定、敏感性分析与 Level 5 视觉复核。

## 结论

本文逐问完成了模型建立、求解、结果解释与可靠性检查。{' '.join(_compact_text(part, maximum=350) for part in abstract_parts)} 所有结论只在 Evidence Pack 固定的接受版本、数据范围和警告边界内成立。

## 参考文献

{chr(10).join(reference_lines)}

## 附录：复现说明、文件清单和核心代码索引

论文由 Evidence Pack `{evidence.get('evidence_hash', '')}` 自动生成。复现时先核验该哈希和 `writer_manifest.json`，再运行论文构建命令；计算代码及结果路径以证据包登记清单为准，正文不重复粘贴完整代码。
"""
    version_dir.mkdir(parents=True, exist_ok=True)
    output = version_dir / "paper.md"
    write_text(output, content)
    manifest = {
        "schema_version": 1,
        "problem_id": problem_id,
        "paper_version": version_dir.name,
        "created_at": utc_now(),
        "generator": "deterministic_evidence_writer",
        "evidence_hash": evidence.get("evidence_hash"),
        "paper_md_sha256": _sha256(output),
        "evidence_ids": sorted(
            str(item.get("evidence_id"))
            for group in ("artifacts", "figures", "citations")
            for item in evidence.get(group, [])
            if item.get("evidence_id")
        ),
        "figure_ids": [str(item.get("stable_id")) for item in evidence.get("figures", [])],
        "citation_ids": citation_keys,
        "warnings_disclosed": warning_lines,
    }
    write_json(version_dir / "writer_manifest.json", manifest)
    return output


def prepare_paper_writing(problem_id: str) -> dict[str, Any]:
    evidence = build_evidence_pack(problem_id)
    version_name, version_dir = create_paper_version(problem_id)
    draft = generate_evidence_markdown(problem_id, version_dir, evidence)
    return {
        "problem_id": problem_id,
        "paper_version": version_name,
        "version_dir": relative(version_dir),
        "draft": relative(draft),
        "evidence": relative(problem_dir(problem_id) / "paper" / "evidence" / "evidence_pack.json"),
        "evidence_hash": evidence["evidence_hash"],
    }


def create_reference_doc(path: Path) -> Path:
    """生成固定的中文数学建模竞赛 Word 样式模板。"""
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    for name, size, bold in (("Normal", 10.5, False), ("Title", 18, True), ("Heading 1", 15, True), ("Heading 2", 13, True), ("Heading 3", 11, True)):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体" if name == "Normal" else "黑体")
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.25
    document.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)
    document.add_heading("数学建模论文样式", 0)
    document.add_paragraph("此段用于固定正文、标题、公式、题注和表格样式。")
    document.add_section(WD_SECTION.NEW_PAGE)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def export_docx_to_pdf(docx_path: Path, pdf_path: Path, timeout_seconds: int) -> None:
    """在隔离 PowerShell 子进程中调用 Microsoft Word 导出 PDF。"""
    script_path = docx_path.parent / "word_export.ps1"
    script = r'''param([string]$Docx, [string]$Pdf)
$ErrorActionPreference = "Stop"
$word = $null
$document = $null
try {
  $word = New-Object -ComObject Word.Application
  $word.Visible = $false
  $word.DisplayAlerts = 0
  $document = $word.Documents.Open([System.IO.Path]::GetFullPath($Docx), $false, $true)
  $document.ExportAsFixedFormat([System.IO.Path]::GetFullPath($Pdf), 17)
} finally {
  if ($null -ne $document) { $document.Close($false); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($document) }
  if ($null -ne $word) { $word.Quit(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($word) }
  [GC]::Collect()
  [GC]::WaitForPendingFinalizers()
}
'''
    write_text(script_path, script)
    command = [
        "powershell.exe",
        "-NoProfile",
        "-NonInteractive",
        "-ExecutionPolicy",
        "Bypass",
        "-File",
        str(script_path),
        "-Docx",
        str(docx_path),
        "-Pdf",
        str(pdf_path),
    ]
    try:
        result = subprocess.run(
            command,
            cwd=docx_path.parent,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout_seconds,
            shell=False,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise TimeoutError(f"Microsoft Word PDF 导出超时（{timeout_seconds} 秒）") from exc
    if result.returncode != 0 or not pdf_path.is_file():
        detail = (result.stderr or result.stdout or "未生成 PDF").strip()[-2000:]
        raise RuntimeError(f"Microsoft Word PDF 导出失败：{detail}")


def inspect_rendered_paper(version_dir: Path, minimum_pdf_pages: int) -> dict[str, Any]:
    """解析 DOCX/PDF，核验媒体、可编辑公式、页数与 A4 页面。"""
    import zipfile

    from pypdf import PdfReader

    docx_path = version_dir / "paper.docx"
    pdf_path = version_dir / "paper.pdf"
    with zipfile.ZipFile(docx_path) as archive:
        document_xml = archive.read("word/document.xml")
        media_count = len([name for name in archive.namelist() if name.startswith("word/media/")])
    reader = PdfReader(str(pdf_path))
    page_count = len(reader.pages)
    if page_count < minimum_pdf_pages:
        raise RuntimeError(f"PDF 页数 {page_count} 低于下限 {minimum_pdf_pages}")
    non_a4_pages: list[int] = []
    blank_pages: list[int] = []
    for index, page in enumerate(reader.pages, 1):
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        short, long = sorted((width, height))
        if abs(short - 595.276) > 12 or abs(long - 841.89) > 12:
            non_a4_pages.append(index)
        if not (page.extract_text() or "").strip():
            blank_pages.append(index)
    consecutive_blank = any(right == left + 1 for left, right in zip(blank_pages, blank_pages[1:]))
    if non_a4_pages:
        raise RuntimeError(f"PDF 存在非 A4 页面：{non_a4_pages}")
    if consecutive_blank:
        raise RuntimeError(f"PDF 存在连续空白页：{blank_pages}")
    return {
        "docx_media_count": media_count,
        "docx_omml_count": document_xml.count(b"<m:oMath"),
        "pdf_pages": page_count,
        "pdf_blank_pages": blank_pages,
        "pdf_non_a4_pages": non_a4_pages,
    }


def render_paper(
    version_dir: Path,
    config: dict[str, Any] | None = None,
    *,
    pdf_exporter: Callable[[Path, Path, int], None] | None = None,
) -> dict[str, Any]:
    """将同一 Markdown 源渲染为 DOCX、TEX，并从 DOCX 导出 PDF。"""
    settings = dict(read_yaml(CONFIG_DIR / "paper.yaml", {}))
    if config:
        settings.update(config)
    markdown = version_dir / "paper.md"
    if not markdown.is_file():
        raise RuntimeError(f"论文 Markdown 不存在：{markdown}")
    pandoc_setting = str(settings.get("pandoc_executable", "pandoc"))
    pandoc = shutil.which(pandoc_setting) or (pandoc_setting if Path(pandoc_setting).is_file() else None)
    if not pandoc:
        raise RuntimeError(f"Pandoc 不可用：{pandoc_setting}")
    render_source = version_dir / "paper.render.md"
    source = markdown.read_text(encoding="utf-8")
    source = _EVIDENCE_MARKER_RE.sub("", source)
    source = _CITATION_RE.sub(lambda match: f"[{match.group(1)}]", source)
    write_text(render_source, source)
    docx_path = version_dir / "paper.docx"
    tex_path = version_dir / "paper.tex"
    pdf_path = version_dir / "paper.pdf"
    resource_paths = [str(version_dir), str(ROOT)]
    base = [str(pandoc), str(render_source), "--from", "markdown+tex_math_dollars", "--resource-path", ";".join(resource_paths)]
    docx_command = [*base, "--to", "docx", "--output", str(docx_path), "--number-sections"]
    reference_value = settings.get("reference_doc")
    if reference_value:
        reference = (ROOT / str(reference_value)).resolve()
        if reference.is_file():
            docx_command.extend(["--reference-doc", str(reference)])
    commands = [docx_command, [*base, "--to", "latex", "--output", str(tex_path), "--number-sections"]]
    pandoc_logs: list[dict[str, Any]] = []
    for command in commands:
        result = subprocess.run(
            command,
            cwd=version_dir,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            shell=False,
            check=False,
        )
        pandoc_logs.append({"command": command, "returncode": result.returncode, "stderr": result.stderr[-4000:]})
        if result.returncode != 0:
            raise RuntimeError(f"Pandoc 转换失败：{result.stderr.strip()[-2000:]}")
    report: dict[str, Any] = {
        "checked_at": utc_now(),
        "paper_version": version_dir.name,
        "source_sha256": _sha256(markdown),
        "docx_sha256": _sha256(docx_path),
        "tex_sha256": _sha256(tex_path),
        "pandoc": pandoc_logs,
        "status": "FAILED_RENDER",
        "errors": [],
    }
    try:
        exporter = pdf_exporter or export_docx_to_pdf
        exporter(docx_path, pdf_path, int(settings.get("pdf_export_timeout_seconds", 120)))
        inspection = inspect_rendered_paper(version_dir, int(settings.get("minimum_pdf_pages", 12)))
        report.update(inspection)
        report["pdf_sha256"] = _sha256(pdf_path)
        report["status"] = "PASS"
    except (OSError, RuntimeError, TimeoutError, ValueError) as exc:
        if pdf_path.exists():
            pdf_path.unlink()
        report["errors"].append(str(exc))
    write_json(version_dir / "render_report.json", report)
    return report
